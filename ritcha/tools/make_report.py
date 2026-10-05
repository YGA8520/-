"""Hebrew report: what was added, which duplicates were removed, what needs the editor's decision."""
import sys, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import merge
from merge import Item, cov, qtext
from model import build

FILES = {
 'z1 סה-סו (גיליון מודפס)': 'ריתחא דאורייתא סימן סה - סו.docx', 'z1 סימן ח': 'ריתחא דאורייתא סימן ח.docx',
 'z2 מסמך חדש (1)': '__מסמך של Microsoft Word חדש (1).docx', 'z2 ריתחא נשא': '__ריתחא - נשא.docx',
 'z2 אהרן לייב פריזנד': 'אהרן לייב פריזנד זאדה.docx', 'z2 זאדה מז': 'זאדה ריתחא - מז.docx', 'z2 סימן ו': 'סימן ו-ז אשר יצר.docx',
 'z2 ריתחא דאורייתא ב': 'ריתחא דאורייתא ב.docx', 'z2 הלכות ציצית': 'ריתחא דאורייתא הלכות ציצית ק.docx',
 'z3 סימן סו': 'ריתחא דאורייתא סימן סו.docx', 'z3 ריתחא (מז)': 'ריתחא דאורייתא.docx', 'z3 ריתחא סימן ח': 'ריתחא סימן ח.docx',
 'z3 ריתחא סימן סד': 'ריתחא סימן סד.docx', 'z3 ריתחא במדבר': 'ריתחא.docx', 'z3 נטילת ידים': 'שאלה הנוטל ידיו שחרית (1).docx',
 'z3 שאלות לחידודא': 'שאלו לחידודא - מז.docx', 'z3 תפילה וקריאת שמע': 'שאלות בעניני תפלה וקריאת שמע.docx',
 'z3 קכג': 'שאלות וחידונים סימן קכג.docx', 'z3 שאלות ה (2)': 'שאלות לסימן ה (2).docx', 'z3 שאלות ה (1)': 'שאלות סימן ה (1).docx',
 'z3 שאלות שונות': 'שאלות שונות בעניני תפלה.docx', 'PDF סימן עא (עא1)': 'ריתחא דאורייתא סימן עא1.pdf (נקרא מהתמונה)',
 'חוברת קיימת': 'החוברת הקודמת',
}
def fname(label):
    return ' ← '.join(FILES.get(x.strip(), x.strip()) for x in label.split('←'))

def setup(doc):
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:cs'), 'Arial')
    sec = doc.sections[0]; sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    for m in ('left_margin', 'right_margin'): setattr(sec, m, Cm(2))
    sec._sectPr.append(OxmlElement('w:bidi'))

def P(doc, text='', size=11, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.RIGHT, before=0, after=4, indent=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after)
    if indent: pf.right_indent = Cm(indent)
    p._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
    p.alignment = align
    if text:
        r = p.add_run(text); r.font.size = Pt(size); r.bold = bold
        r._r.get_or_add_rPr().append(OxmlElement('w:rtl'))
        if color: r.font.color.rgb = RGBColor.from_string(color)
    return p

def H(doc, text): return P(doc, text, size=14, bold=True, color='404040', before=12, after=4)
def B(doc, text, indent=0.6): return P(doc, '• ' + text, indent=indent)

def main(out):
    part1, part2, log, pool = merge.merge()
    comp = build('input_document.xml')
    cut = next(i for i, s in enumerate(comp) if s['title'] and s['title'].startswith('שאלה להלכה למעשה'))
    old = [Item(s['num'], e['blocks'], '') for s in comp[:cut] for e in s['entries'] if e['type'] == 'item']
    n_old = len(old)
    rows = []        # (siman, kind, snippet, item)
    for s in part1:
        name = ('סימן ' + s['num']) if s['num'] else s['title']
        for e in s['entries']:
            F = Item(s['num'], e['blocks'], '')
            matched_old = max(max(cov(F, O), cov(O, F)) for O in old) >= 0.6
            replaced = any('ניסוח קודם' in d for d in e['dups'])
            kind = 'edited' if replaced else ('old' if matched_old else 'new')
            rows.append((name, kind, qtext(e['blocks']), e))
    n_new = sum(1 for r in rows if r[1] == 'new'); n_edit = sum(1 for r in rows if r[1] == 'edited')
    n_dup = sum(len(r[3]['dups']) for r in rows) - n_edit
    n2 = sum(1 for s in part2 for e in s['entries'] if e['type'] == 'item')

    doc = Document(); setup(doc)
    P(doc, 'ריתחא דאורייתא – דוח שילוב שאלות חדשות', size=18, bold=True, color='404040', align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

    H(doc, 'בקצרה')
    B(doc, 'החלק הראשון של החוברת (שאלות לפי סימנים) מכיל עכשיו %d שאלות במקום %d.' % (len(rows), n_old))
    B(doc, '%d שאלות חדשות נוספו (מהקבצים החדשים, מתוכן 9 שנקראו ידנית מתמונה ב-PDF).' % n_new)
    B(doc, '%d שאלות הופיעו כבר בחוברת בנוסח מוקדם והוחלפו בנוסח המעובד של הגיליון המודפס (פירוט למטה).' % n_edit)
    B(doc, '%d כפילויות (אותה שאלה בקובץ נוסף) זוהו ונשארה רק הפעם הראשונה.' % n_dup)
    B(doc, 'חלק "שאלה להלכה למעשה" (%d שאלות) נשאר כמות שהוא (לא נמצאו בו כפילויות, לא עם שאר החוברת ולא בתוכו).' % n2)

    H(doc, 'עיצוב – מה שונה בגרסה זו')
    for t in ['כותרת העמוד היא שורה אחת: "סימן X" משמאל, העיטור כמפריד באמצע, ו"ריתחא דאורייתא" מימין. הכותרת קטנה יותר מקודם וציון הסימן גדול יותר.',
              'הגופנים: טקסט רץ – Tehila (16); כותרת סימן ואותיות המספור – RimonMF; שמות השואלים והמקורות – Gisha; מספור עמודים – Times New Roman.',
              'כל אות לפני שאלה ממורכזת בדיוק בין שני הריבועים שמצדדיה – גם אנכית וגם אופקית (המיקום חושב לפי צורת האות בגופן, לכל אות בנפרד).',
              'בלוק הציטוטים והמקורות שבסימן מז (בקובץ "זאדה ריתחא – מז") הוסר לגמרי.',
              'נשארו בתוקף: אין כותרות משנה, אין לוגו ו"ברצות ה\'", הפתיחה של "שאלה להלכה למעשה" בעמוד חדש כשעיטור הסיום בעמוד שלפניה, והסרת הנוסח "הלכה למעשה:" מתחילת השאלות.']:
        B(doc, t)

    H(doc, 'שאלות חדשות שנוספו')
    cur = None
    for name, kind, text, e in rows:
        if kind != 'new': continue
        if name != cur: P(doc, name, bold=True, before=6, after=2); cur = name
        B(doc, '%s…  (%s)' % (text[:70], fname(e['src'].split('←')[0].strip() if 'חוברת קיימת' not in e['src'].split('←')[0] else e['src'].split('←')[-1].strip())), indent=1.2)

    H(doc, 'שאלות שהוחלפו בנוסח המעובד')
    for name, kind, text, e in rows:
        if kind != 'edited': continue
        B(doc, '%s – %s… (נשאר הנוסח מ-%s; הנוסח המוקדם הוסר)' % (name, text[:60], fname(e['src'].split('←')[0].strip())))
    P(doc, 'אם תעדיף את הנוסח המוקדם (הארוך) – אפשר להחזיר.', size=10, color='666666', indent=0.6)

    H(doc, 'כפילויות שהוסרו (לפי סימן)')
    cur = None
    for name, kind, text, e in rows:
        if not e['dups'] or kind == 'edited': continue
        if name != cur: P(doc, name, bold=True, before=6, after=2); cur = name
        B(doc, '%s…  –  מופיעה גם ב: %s' % (text[:55], '; '.join(sorted({fname(d) for d in e['dups']}))), indent=1.2)

    H(doc, 'נקודות לבדיקה / החלטה')
    for t in ['הטקסט שב-RimonMF (כותרות הסימנים, אותיות המספור וציון הסימן בראש העמוד) שמור בקידוד של הגופן – כמו Wingdings – כדי שיוצג בדיוק כמו ב-PDF. להוספת סימן חדש הכי נוח להעתיק כותרת סימן קיימת ולשנות אותה; אם בוורד שלך הקלדה עברית ישירה בגופן הזה עובדת – אפשר גם כך.',
              'שורה "שאלה לענות אמן באמצע ק"ש" (אהרן לייב פריזנד) מופיעה בקובץ של סימן ח, אבל מצוין בה "סי\' סו סעיף ג" – שיבצתי אותה בסימן סו.',
              'שאלות על קברים בסימן עא (9 שאלות) נמצאו רק ב-PDF "סימן עא1" ללא קובץ וורד – הקלדתי אותן ידנית מהתמונה; כדאי לקרוא שוב.',
              'בחלק "הלכה למעשה" נשארה השורה "סימן נה סעיף ב" בין סימן נו לסימן נו – כנראה טעות הקלדה (צ"ל נו). לא תיקנתי.',
              'השמות בסוף השאלות הוחלפו לנוסח שבגיליונים המודפסים: "הרב X" ← "[הערת ה"ר X]". "בני החבורה" נשאר.',
              'בעמוד שיש בו כמה סימנים קטנים מוצג בכותרת הסימן הראשון בעמוד. אפשר להציג טווח (לדוגמה "סימן נא – נב") – ידרוש שדה מותנה שפועל רק בוורד.',
              'שאלות שהובאו בנוסח ארוך וגם בנוסח קצר (סימן סה ×3, סימן סד) – נשאר הנוסח המעובד שבגיליון.']:
        B(doc, t)
    doc.save(out)
    print('report', out, '| new', n_new, 'edited', n_edit, 'dups', n_dup, 'total part1', len(rows), 'old', n_old)

if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'report.docx')

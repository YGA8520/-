#!/usr/bin/env python3
"""Small helper: a Hebrew report made of sections with tables -> PDF (Chromium) and DOCX (python-docx), both right-to-left."""
import os, subprocess
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

HERE = os.path.dirname(os.path.abspath(__file__))
esc = lambda s: str(s).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def write_report(outdir, base, title, intro, sections, widths=None):
    """sections: [(heading, [column names], [[cell, ...], ...], optional note text)]"""
    html = ('<html dir="rtl" lang="he"><head><meta charset="utf-8"><style>'
            '@font-face{font-family:D;src:url("assets/fonts/DavidLibre-Regular.ttf")}@font-face{font-family:D;font-weight:700;src:url("assets/fonts/DavidLibre-Bold.ttf")}'
            '@page{size:A4;margin:14mm}body{font:10.5pt/1.45 D;color:#111}h1{font-size:20pt;margin:0 0 3mm}h2{font-size:13.5pt;margin:7mm 0 2mm}p{margin:0 0 3mm}'
            'table{width:100%;border-collapse:collapse;margin-bottom:3mm}th,td{border:0.6pt solid #888;padding:1.2mm 2mm;text-align:right;vertical-align:top;font-size:9.5pt}'
            'th{background:#eee}tr{page-break-inside:avoid}</style></head><body>' + f'<h1>{esc(title)}</h1>' + ''.join(f'<p>{esc(x)}</p>' for x in intro))
    for sec in sections:
        head, cols, rows = sec[0], sec[1], sec[2]
        html += f'<h2>{esc(head)}</h2>'
        if len(sec) > 3 and sec[3]:
            html += f'<p>{esc(sec[3])}</p>'
        html += '<table><tr>' + ''.join(f'<th>{esc(c)}</th>' for c in cols) + '</tr>' + ''.join('<tr>' + ''.join(f'<td>{esc(c)}</td>' for c in r) + '</tr>' for r in rows) + '</table>'
    html += '</body></html>'
    hp = os.path.join(HERE, '_report.html')
    open(hp, 'w').write(html)
    js = ("let chromium;try{({chromium}=require('playwright'))}catch(e){({chromium}=require('/opt/node22/lib/node_modules/playwright'))}(async()=>{const b=await chromium.launch();const p=await b.newPage();"
          f"await p.goto('file://{hp}');await p.evaluate(()=>document.fonts.ready);await p.pdf({{path:'{outdir}/{base}.pdf',format:'A4',printBackground:true,preferCSSPageSize:true}});await b.close();}})();")
    subprocess.run(['node', '-e', js], check=True)
    os.remove(hp)

    def rtl(par):
        par._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def run(par, text, bold=False, size=10):
        r = par.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        r.font.name = 'David'
        rpr = r._r.get_or_add_rPr()
        rpr.append(OxmlElement('w:rtl'))
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.insert(0, rf)
        rf.set(qn('w:cs'), 'David')

    wd = docx.Document()
    sec0 = wd.sections[0]
    sec0.left_margin = sec0.right_margin = Cm(1.5)
    p = wd.add_paragraph(); rtl(p); run(p, title, True, 18)
    for x in intro:
        p = wd.add_paragraph(); rtl(p); run(p, x, False, 11)
    for sec in sections:
        head, cols, rows = sec[0], sec[1], sec[2]
        p = wd.add_paragraph(); rtl(p); run(p, head, True, 13)
        if len(sec) > 3 and sec[3]:
            p = wd.add_paragraph(); rtl(p); run(p, sec[3], False, 10)
        t = wd.add_table(rows=1, cols=len(cols)); t.style = 'Table Grid'
        t._tbl.tblPr.append(OxmlElement('w:bidiVisual'))
        for i, c in enumerate(cols):
            cell = t.rows[0].cells[i]; cell.text = ''
            q = cell.paragraphs[0]; rtl(q); run(q, c, True, 9.5)
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = ''
                q = cells[i].paragraphs[0]; rtl(q); run(q, str(v), False, 9.5)
    wd.save(os.path.join(outdir, base + '.docx'))

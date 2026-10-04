#!/usr/bin/env python3
"""Notes about irregularities in the SOURCE text, as PDF and Word (easy to download / forward).
usage: make_notes.py book.doc.json out/book.layout.json outdir"""
import json, os, re, subprocess, sys
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

doc_path, lay_path, outdir = sys.argv[1:4]
d = json.load(open(doc_path)); lay = json.load(open(lay_path))
T = lambda runs: ''.join(r['t'] for r in runs if 'fn' not in r)
rows = []


def snip(t, a, b, ctx=34):
    """excerpt around t[a:b], trimmed to whole words, with … where the text goes on"""
    lo, hi = max(0, a - ctx), min(len(t), b + ctx)
    while lo > 0 and lo < a and not t[lo - 1].isspace():
        lo += 1
    while hi < len(t) and hi > b and not t[hi].isspace():
        hi -= 1
    return ('… ' if lo > 0 else '') + t[lo:hi].strip() + (' …' if hi < len(t) else '')


for ai, a in enumerate(d['articles']):
    pg = lay['toc'][ai]['page']
    for b in a['blocks']:
        if b['t'] != 'p':
            continue
        t = T(b['runs']).strip()
        for o, c, nm in (('(', ')', 'סוגריים עגולים ( ) לא מאוזנים'), ('[', ']', 'סוגריים מרובעים [ ] לא מאוזנים')):
            if t.count(o) != t.count(c):
                bal, pos = 0, None
                for i, ch in enumerate(t):
                    if ch == o:
                        bal += 1
                        if bal == 1:
                            pos = i
                    elif ch == c:
                        bal -= 1
                        if bal < 0:
                            pos, bal = i, 0
                            break
                if pos is None:
                    pos = max(t.find(o), t.find(c), 0)
                rows.append((pg, a['title'], nm, snip(t, pos, pos + 1)))
        m = re.search(r'\s[,.;:](?![\d.])', t)
        if m and not re.search(r'\s[,.;:]$', t):
            rows.append((pg, a['title'], 'רווח לפני סימן פיסוק', snip(t, m.start(), m.end())))
        if '???' in t:
            i = t.index('???')
            rows.append((pg, a['title'], 'סימני שאלה (???) בטקסט', snip(t, i, i + 3)))
title = 'הערות על הטקסט המקורי'
intro = ('בקריאה המלאה של החוברת נמצאו מקומות שבהם הטקסט עצמו אינו תקין, כגון סוגריים שנפתחו ולא נסגרו. '
         'רווחים שהוקלדו לפני סימני פיסוק תוקנו בחוברת. סוגריים לא מאוזנים לא תוקנו (להתעלם בשלב זה) ולא שיניתי תוכן. '
         'מי שרוצה יתקן במקור וישלח שוב, והחוברת תיבנה מחדש. מספרי העמודים הם לפי מספור החוברת.')
# ---------------- PDF (via Chromium) ----------------
esc = lambda s: s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
html = ('<html dir="rtl" lang="he"><head><meta charset="utf-8"><style>'
        '@font-face{font-family:D;src:url("assets/fonts/DavidLibre-Regular.ttf")}@font-face{font-family:D;font-weight:700;src:url("assets/fonts/DavidLibre-Bold.ttf")}'
        '@page{size:A4;margin:16mm}body{font:11pt/1.5 D;color:#111}h1{font-size:20pt;margin:0 0 4mm}p{margin:0 0 6mm}'
        'table{width:100%;border-collapse:collapse}th,td{border:0.6pt solid #888;padding:2mm 2.5mm;text-align:right;vertical-align:top;font-size:10pt}'
        'th{background:#eee}td.pg{width:14mm;text-align:center}td.ty{width:38mm}</style></head><body>'
        f'<h1>{title}</h1><p>{intro}</p><table><tr><th>עמוד</th><th>מאמר</th><th>סוג</th><th>קטע</th></tr>' +
        ''.join(f'<tr><td class="pg">{p}</td><td>{esc(t)}</td><td class="ty">{esc(k)}</td><td>{esc(x)}</td></tr>' for p, t, k, x in rows) + '</table></body></html>')
hp = os.path.join(os.path.dirname(os.path.abspath(__file__)), '_notes.html')
open(hp, 'w').write(html)
js = ("let chromium;try{({chromium}=require('playwright'))}catch(e){({chromium}=require('/opt/node22/lib/node_modules/playwright'))}(async()=>{const b=await chromium.launch();const p=await b.newPage();"
      f"await p.goto('file://{hp}');await p.evaluate(()=>document.fonts.ready);await p.pdf({{path:'{outdir}/source-notes.pdf',format:'A4',printBackground:true,preferCSSPageSize:true}});await b.close();}})();")
subprocess.run(['node', '-e', js], check=True)
os.remove(hp)
# ---------------- Word ----------------
def rtl(par):
    ppr = par._p.get_or_add_pPr()
    ppr.append(OxmlElement('w:bidi'))
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
def put(cell, text, bold=False, size=10):
    cell.text = ''
    par = cell.paragraphs[0]
    rtl(par)
    r = par.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = 'David'
    rpr = r._r.get_or_add_rPr()
    rtl_el = OxmlElement('w:rtl'); rpr.append(rtl_el)
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    rf.set(qn('w:cs'), 'David')
wd = docx.Document()
sec = wd.sections[0]
sec.left_margin = sec.right_margin = Cm(1.8)
h = wd.add_paragraph(); rtl(h); hr = h.add_run(title); hr.bold = True; hr.font.size = Pt(18); hr.font.name = 'David'
hr._r.get_or_add_rPr().append(OxmlElement('w:rtl'))
p = wd.add_paragraph(); rtl(p); pr = p.add_run(intro); pr.font.size = Pt(11); pr.font.name = 'David'; pr._r.get_or_add_rPr().append(OxmlElement('w:rtl'))
tbl = wd.add_table(rows=1, cols=4); tbl.style = 'Table Grid'
tblPr = tbl._tbl.tblPr; tblPr.append(OxmlElement('w:bidiVisual'))
for i, t in enumerate(('עמוד', 'מאמר', 'סוג', 'קטע')):
    put(tbl.rows[0].cells[i], t, True)
for p_, t, k, x in rows:
    cells = tbl.add_row().cells
    for i, v in enumerate((str(p_), t, k, x)):
        put(cells[i], v)
for row in tbl.rows:
    for i, w in enumerate((Cm(1.5), Cm(5), Cm(4), Cm(6.7))):
        row.cells[i].width = w
wd.save(os.path.join(outdir, 'source-notes.docx'))
print(len(rows), 'items ->', outdir)

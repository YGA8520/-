#!/usr/bin/env python3
"""Turn the pilot markdown files into one right-to-left Hebrew Word file.

Usage: md_to_docx_rtl.py OUT.docx "Title" intro.md file1.md file2.md ...  (each file starts on a new page)
Supported markdown: # / ## headings, paragraphs, - bullets, 1. numbers, | tables |, ---, **bold**, *italic*, and the
⟦...⟧ working markers, which are shown small and grey.
"""
import re, sys
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

FONT = 'David'
TOKEN = re.compile(r'(⟦.*?⟧|\*\*.+?\*\*|\*.+?\*)')


def rtl_par(par):
    pPr = par._p.get_or_add_pPr()
    pPr.append(OxmlElement('w:bidi'))


def style_run(run, size=13, bold=None, italic=None, color=None):
    rPr = run._r.get_or_add_rPr()
    fonts = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        fonts.set(qn(a), FONT)
    rPr.append(fonts)
    rPr.append(OxmlElement('w:rtl'))
    run.font.size = Pt(size)
    szcs = OxmlElement('w:szCs'); szcs.set(qn('w:val'), str(size * 2)); rPr.append(szcs)
    if bold:
        run.bold = True
        rPr.append(OxmlElement('w:bCs'))
    if italic:
        run.italic = True
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_runs(par, text, size=13, bold=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith('⟦'):
            style_run(par.add_run(part), size=size - 3, color='7A7A7A')
        elif part.startswith('**') and part.endswith('**') and len(part) > 4:
            style_run(par.add_run(part[2:-2]), size=size, bold=True)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            style_run(par.add_run(part[1:-1]), size=size, italic=True, bold=bold)
        else:
            style_run(par.add_run(part), size=size, bold=bold)


def new_par(doc, text='', size=13, bold=False, space_after=6, indent_cm=0, hanging_cm=0, keep_next=False):
    par = doc.add_paragraph()
    rtl_par(par)
    pf = par.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.25
    if indent_cm:
        pf.right_indent = Cm(indent_cm)
    if hanging_cm:
        pf.first_line_indent = Cm(-hanging_cm)
    if keep_next:
        pf.keep_with_next = True
    add_runs(par, text, size=size, bold=bold)
    return par


def add_table(doc, rows):
    ncols = len(rows[0])
    table = doc.add_table(rows=0, cols=ncols)
    table.style = 'Table Grid'
    tblPr = table._tbl.tblPr
    tblPr.append(OxmlElement('w:bidiVisual'))
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, txt in enumerate(row):
            par = cells[j].paragraphs[0]
            rtl_par(par)
            add_runs(par, txt.strip(), size=11, bold=(i == 0))
    doc.add_paragraph()


def render_md(doc, path):
    lines = open(path, encoding='utf-8').read().split('\n')
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln:
            i += 1
        elif ln.startswith('# '):
            new_par(doc, ln[2:], size=20, bold=True, space_after=10, keep_next=True)
            i += 1
        elif ln.startswith('## '):
            new_par(doc, ln[3:], size=15, bold=True, space_after=4, keep_next=True)
            i += 1
        elif ln.strip() == '---':
            i += 1
        elif ln.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'\s*-+\s*', c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
        elif ln.startswith('- ') or ln.startswith('  - '):
            sub = ln.startswith('  ')
            new_par(doc, '• ' + ln.strip()[2:], indent_cm=1.4 if sub else 0.6, hanging_cm=0.5, space_after=3)
            i += 1
        elif re.match(r'\d+\. ', ln):
            n, rest = ln.split('. ', 1)
            new_par(doc, n + '. ' + rest, indent_cm=0.6, hanging_cm=0.6, space_after=3)
            i += 1
        else:
            new_par(doc, ln)
            i += 1


def main():
    out, title, intro, *files = sys.argv[1:]
    doc = Document()
    sec = doc.sections[0]
    sec.right_margin = sec.left_margin = Cm(2.4)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    bidi = OxmlElement('w:bidi')
    sec._sectPr.append(bidi)
    new_par(doc, title, size=24, bold=True, space_after=14)
    render_md(doc, intro)
    for f in files:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        render_md(doc, f)
    doc.save(out)
    print('wrote', out)


if __name__ == '__main__':
    main()
